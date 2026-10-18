# -*- coding: utf-8 -*-
"""
GUI Qt6 (PySide6) cho SoilData2AI (lõi xlsx2csv.py).

Chạy trực tiếp:
    python src/gui.py
Đóng gói:
    build_exe.ps1  (hoặc lệnh PyInstaller trong README)
"""
from __future__ import annotations

import argparse
import contextlib
import csv
import html
import io
import os
import sys
import tempfile
import time
from datetime import datetime
from pathlib import Path

from PySide6.QtCore import QSettings, QSize, Qt, QThread, QTimer, QUrl, Signal
from PySide6.QtGui import QColor, QDesktopServices, QFont, QIcon, QPixmap
from PySide6.QtWidgets import (
    QAbstractItemView,
    QApplication,
    QCheckBox,
    QFileDialog,
    QFrame,
    QGridLayout,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QProgressBar,
    QPushButton,
    QSizePolicy,
    QTableWidget,
    QTableWidgetItem,
    QTabWidget,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)

import xlsx2csv as core

APP_NAME = "SoilData2AI"
APP_VERSION = "2.0"
APP_ORG = "MDC"


def project_root() -> Path:
    """Thư mục gốc khi chạy source; thư mục giải nén khi chạy EXE onefile."""
    if getattr(sys, "frozen", False):
        return Path(getattr(sys, "_MEIPASS", Path(sys.executable).resolve().parent))
    return Path(__file__).resolve().parent.parent


def resource_path(rel: str) -> Path:
    return project_root() / rel


def default_dirs() -> tuple[Path, Path]:
    here = Path(__file__).resolve().parent
    candidates = [
        here.parent / "dataXLSX",
        here.parent.parent / "dataXLSX",
        Path.cwd() / "dataXLSX",
    ]
    for p in candidates:
        if p.is_dir():
            out = p.parent / "dataCSV_UTF8"
            return p, out
    home = Path.home()
    return home / "dataXLSX", home / "dataCSV_UTF8"


class FolderLineEdit(QLineEdit):
    folderDropped = Signal(str)

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.setAcceptDrops(True)

    def dragEnterEvent(self, e):  # noqa: N802
        if e.mimeData().hasUrls():
            for u in e.mimeData().urls():
                if u.isLocalFile() and Path(u.toLocalFile()).is_dir():
                    e.acceptProposedAction()
                    return
        super().dragEnterEvent(e)

    def dropEvent(self, e):  # noqa: N802
        for u in e.mimeData().urls():
            p = Path(u.toLocalFile())
            if p.is_dir():
                self.setText(str(p))
                self.folderDropped.emit(str(p))
                e.acceptProposedAction()
                return
        super().dropEvent(e)


class EmitStream(io.TextIOBase):
    """Chuyển print()/stderr của lõi thành tín hiệu log của GUI."""

    def __init__(self, emit):
        super().__init__()
        self._emit = emit
        self._buf = ""

    def write(self, s: str) -> int:  # type: ignore[override]
        self._buf += s
        while "\n" in self._buf:
            line, self._buf = self._buf.split("\n", 1)
            if line.strip():
                self._emit(line.rstrip())
        return len(s)

    def flush(self) -> None:  # type: ignore[override]
        if self._buf.strip():
            self._emit(self._buf.strip())
            self._buf = ""


class ConvertWorker(QThread):
    log = Signal(str, str)                 # level, message
    progress = Signal(int, int, str)       # current, total, filename
    file_done = Signal(str, int, str)      # filename, sheets, error
    done = Signal(dict)

    def __init__(
        self,
        in_dir: Path,
        out_dir: Path,
        merge_alldata: bool = False,
        parent=None,
    ):
        super().__init__(parent)
        self.in_dir = in_dir
        self.out_dir = out_dir
        self.merge_alldata = merge_alldata
        self._stop = False

    def request_stop(self) -> None:
        self._stop = True

    def run(self) -> None:  # noqa: D401
        started = time.perf_counter()
        stats = {
            "files": 0,
            "ok": 0,
            "failed": 0,
            "sheets": 0,
            "elapsed": 0.0,
            "stopped": False,
            "merged": False,
            "output": str(self.out_dir),
        }
        try:
            files = core.discover_workbooks(self.in_dir)
            stats["files"] = len(files)
            if not files:
                self.log.emit("WARN", "Không tìm thấy tệp .xls/.xlsx trong thư mục đầu vào.")
                stats["elapsed"] = time.perf_counter() - started
                self.done.emit(stats)
                return

            self.out_dir.mkdir(parents=True, exist_ok=True)
            used: set[str] = set()
            self.log.emit("INFO", f"Bắt đầu: {len(files)} tệp → {self.out_dir}")

            stream = EmitStream(
                lambda line: self.log.emit(
                    "WARN" if ("WARN" in line or "FAIL" in line) else "INFO", line
                )
            )
            with contextlib.redirect_stdout(stream), contextlib.redirect_stderr(stream):
                for i, f in enumerate(files, 1):
                    if self._stop:
                        stats["stopped"] = True
                        self.log.emit("WARN", "Đã dừng theo yêu cầu. Các tệp còn lại được giữ nguyên.")
                        break

                    self.progress.emit(i - 1, len(files), f.name)
                    try:
                        if f.suffix.lower() == ".xlsx":
                            n = core.convert_xlsx(f, self.out_dir, used)
                        else:
                            n = core.convert_xls(f, self.out_dir, used)
                        stats["ok"] += 1
                        stats["sheets"] += n
                        self.log.emit("OK", f"OK   {f.name}   ({n} sheet)")
                        self.file_done.emit(f.name, n, "")
                    except Exception as e:  # noqa: BLE001
                        stats["failed"] += 1
                        msg = str(e).replace("\n", " ")
                        self.log.emit("FAIL", f"FAIL {f.name} → {msg}")
                        self.file_done.emit(f.name, 0, msg)
                    self.progress.emit(i, len(files), f.name)
            stream.flush()

            if self.merge_alldata and stats["ok"] and not self._stop:
                self.log.emit("INFO", "Đang gộp AllData.csv từ các CSV vừa tạo ...")
                try:
                    import normalize

                    out_all = self.out_dir / "AllData.csv"
                    rc = normalize.main(
                        [
                            "--input",
                            str(self.out_dir),
                            "--max-files",
                            "0",
                            "--output",
                            str(out_all),
                        ]
                    )
                    stats["merged"] = rc == 0 and out_all.exists()
                    if stats["merged"]:
                        self.log.emit("OK", f"OK   AllData.csv ({out_all.stat().st_size / 1024 / 1024:.1f} MB)")
                    else:
                        self.log.emit("WARN", "Không gộp được AllData.csv.")
                except Exception as e:  # noqa: BLE001
                    self.log.emit("WARN", f"Lỗi gộp AllData.csv: {e}")

        except Exception as e:  # noqa: BLE001
            stats["failed"] += 1
            self.log.emit("FAIL", f"Lỗi: {e}")
        finally:
            stats["elapsed"] = time.perf_counter() - started
            self.done.emit(stats)


class StatCard(QFrame):
    def __init__(self, title: str, value: str, accent: str, parent=None):
        super().__init__(parent)
        self.setProperty("card", True)
        self.setMinimumHeight(88)
        lay = QVBoxLayout(self)
        lay.setContentsMargins(16, 12, 16, 10)
        lay.setSpacing(2)
        self.value = QLabel(value)
        self.value.setStyleSheet(f"font-size:26px; font-weight:700; color:{accent};")
        cap = QLabel(title.upper())
        cap.setStyleSheet("font-size:9px; font-weight:700; color:#526476; letter-spacing:1px;")
        lay.addWidget(self.value)
        lay.addWidget(cap)
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)

    def set_value(self, v: str) -> None:
        self.value.setText(v)


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(f"{APP_NAME}  ·  Excel → CSV UTF-8")
        self.setMinimumSize(980, 720)
        self.resize(1120, 840)
        icon_path = resource_path("assets/app.ico")
        if icon_path.exists():
            self.setWindowIcon(QIcon(str(icon_path)))

        self.settings = QSettings(APP_ORG, "SoilData2AI")
        self.worker: ConvertWorker | None = None
        self._log_lines: list[tuple[str, str, str]] = []
        self._last_output: Path | None = None

        self._build_ui()
        self._load_settings()

    # ------------------------------------------------------------------ UI
    def _build_ui(self) -> None:
        central = QWidget()
        central.setObjectName("Central")
        self.setCentralWidget(central)
        root = QVBoxLayout(central)
        root.setContentsMargins(22, 18, 22, 12)
        root.setSpacing(12)

        # Header
        header = QHBoxLayout()
        header.setSpacing(14)
        logo = QLabel()
        p = resource_path("assets/app.png")
        if p.exists():
            logo.setPixmap(QPixmap(str(p)).scaled(52, 52, Qt.KeepAspectRatio, Qt.SmoothTransformation))
        header.addWidget(logo)
        titles = QVBoxLayout()
        titles.setSpacing(1)
        h1 = QLabel("Excel → CSV  ·  Chuẩn hóa TCVN3")
        h1.setStyleSheet("font-size:22px; font-weight:700; color:#0B2C4D;")
        sub = QLabel("Chuyển bảng tính .xls/.xlsx sang CSV UTF-8, sẵn sàng cho pipeline AI")
        sub.setStyleSheet("font-size:11px; color:#526476;")
        titles.addWidget(h1)
        titles.addWidget(sub)
        header.addLayout(titles, 1)
        version = QLabel(f"v{APP_VERSION}")
        version.setStyleSheet("font-size:10px; color:#8593A1; background:#FFFFFF; border:1px solid #D7E3EF; border-radius:9px; padding:3px 9px;")
        header.addWidget(version, 0, Qt.AlignTop)
        root.addLayout(header)

        # Folder card
        card = QFrame()
        card.setProperty("card", True)
        grid = QGridLayout(card)
        grid.setContentsMargins(16, 14, 16, 14)
        grid.setHorizontalSpacing(10)
        grid.setVerticalSpacing(9)
        lab_in = QLabel("Thư mục đầu vào")
        lab_in.setStyleSheet("font-size:11px; font-weight:600;")
        self.in_edit = FolderLineEdit()
        self.in_edit.setPlaceholderText("Chọn thư mục chứa .xls / .xlsx …")
        self.in_btn = QPushButton("Chọn…")
        self.in_btn.setObjectName("Secondary")
        self.in_btn.clicked.connect(self._pick_input)
        lab_out = QLabel("Thư mục đầu ra")
        lab_out.setStyleSheet("font-size:11px; font-weight:600;")
        self.out_edit = FolderLineEdit()
        self.out_edit.setPlaceholderText("Chọn thư mục ghi CSV …")
        self.out_btn = QPushButton("Chọn…")
        self.out_btn.setObjectName("Secondary")
        self.out_btn.clicked.connect(self._pick_output)
        grid.addWidget(lab_in, 0, 0)
        grid.addWidget(self.in_edit, 0, 1)
        grid.addWidget(self.in_btn, 0, 2)
        grid.addWidget(lab_out, 1, 0)
        grid.addWidget(self.out_edit, 1, 1)
        grid.addWidget(self.out_btn, 1, 2)
        grid.setColumnStretch(1, 1)

        opts = QHBoxLayout()
        opts.setSpacing(18)
        self.cb_open = QCheckBox("Mở thư mục khi hoàn tất")
        self.cb_open.setChecked(True)
        self.cb_log = QCheckBox("Ghi nhật ký ra tệp")
        self.cb_log.setChecked(True)
        self.cb_merge = QCheckBox("Gộp AllData.csv sau khi chuyển")
        self.cb_merge.setToolTip("Dùng normalize.py để gộp các CSV ngang thành một bảng AllData.csv.")
        for w in (self.cb_open, self.cb_log, self.cb_merge):
            opts.addWidget(w)
        opts.addStretch(1)
        grid.addLayout(opts, 2, 1, 1, 2)
        root.addWidget(card)

        # Actions
        actions = QHBoxLayout()
        actions.setSpacing(10)
        self.start_btn = QPushButton("Bắt đầu chuyển đổi")
        self.start_btn.setObjectName("Primary")
        self.start_btn.setMinimumHeight(40)
        self.start_btn.setMinimumWidth(180)
        self.start_btn.clicked.connect(self._start)
        self.stop_btn = QPushButton("Dừng")
        self.stop_btn.setObjectName("Danger")
        self.stop_btn.setMinimumHeight(40)
        self.stop_btn.setEnabled(False)
        self.stop_btn.clicked.connect(self._stop)
        self.open_btn = QPushButton("Mở thư mục")
        self.open_btn.setObjectName("Secondary")
        self.open_btn.setMinimumHeight(40)
        self.open_btn.clicked.connect(self._open_output)
        self.clear_btn = QPushButton("Xoá nhật ký")
        self.clear_btn.setObjectName("Ghost")
        self.clear_btn.setMinimumHeight(40)
        self.clear_btn.clicked.connect(self._clear_log)
        actions.addWidget(self.start_btn)
        actions.addWidget(self.stop_btn)
        actions.addWidget(self.open_btn)
        actions.addWidget(self.clear_btn)
        actions.addStretch(1)
        self.state_lbl = QLabel("Sẵn sàng")
        self.state_lbl.setStyleSheet("font-size:11px; color:#526476;")
        actions.addWidget(self.state_lbl)
        root.addLayout(actions)

        # Stats
        stats = QHBoxLayout()
        stats.setSpacing(10)
        self.card_files = StatCard("Tệp đã xử lý", "0 / 0", "#1769AA")
        self.card_sheets = StatCard("Sheet → CSV", "0", "#5FA83C")
        self.card_failed = StatCard("Tệp lỗi", "0", "#E8891C")
        self.card_time = StatCard("Thời gian", "0:00", "#0B2C4D")
        for w in (self.card_files, self.card_sheets, self.card_failed, self.card_time):
            stats.addWidget(w)
        root.addLayout(stats)

        # Progress
        pbox = QVBoxLayout()
        pbox.setSpacing(4)
        self.progress = QProgressBar()
        self.progress.setRange(0, 100)
        self.progress.setValue(0)
        self.progress.setTextVisible(False)
        self.progress.setFixedHeight(10)
        self.progress_lbl = QLabel("Chưa bắt đầu")
        self.progress_lbl.setStyleSheet("font-size:10px; color:#526476;")
        pbox.addWidget(self.progress)
        pbox.addWidget(self.progress_lbl)
        root.addLayout(pbox)

        # Tabs
        self.tabs = QTabWidget()
        self.log_view = QTextEdit()
        self.log_view.setReadOnly(True)
        self.log_view.setStyleSheet(
            "QTextEdit { background:#FFFFFF; color:#0B2C4D; border:none; "
            "font-family: Consolas, 'Courier New'; font-size:12px; padding:8px; }"
        )
        self.log_view.document().setMaximumBlockCount(3000)
        self.tabs.addTab(self.log_view, "Nhật ký")
        self.table = QTableWidget(0, 4)
        self.table.setHorizontalHeaderLabels(["Tệp", "Sheet", "Trạng thái", "Ghi chú"])
        self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.ResizeToContents)
        self.table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeMode.ResizeToContents)
        self.table.horizontalHeader().setSectionResizeMode(3, QHeaderView.ResizeMode.Stretch)
        self.table.setAlternatingRowColors(True)
        self.table.verticalHeader().setVisible(False)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.tabs.addTab(self.table, "Kết quả")
        root.addWidget(self.tabs, 1)

        footer = QLabel(
            "Mẹo: kéo–thả thư mục vào ô đường dẫn.  ·  "
            "CSV ghi bằng UTF-8 có BOM; sheet rỗng tự động bỏ qua."
        )
        footer.setStyleSheet("font-size:9px; color:#8593A1;")
        root.addWidget(footer)
        self._append_log(
            "INFO",
            "Sẵn sàng. Chọn thư mục đầu vào/đầu ra rồi bấm “Bắt đầu chuyển đổi”.",
        )

    # -------------------------------------------------------------- helpers
    def _load_settings(self) -> None:
        din, dout = default_dirs()
        self.in_edit.setText(self.settings.value("input", str(din), str))
        self.out_edit.setText(self.settings.value("output", str(dout), str))
        self.cb_open.setChecked(self.settings.value("open", True, bool))
        self.cb_log.setChecked(self.settings.value("log", True, bool))
        self.cb_merge.setChecked(self.settings.value("merge", False, bool))

    def _save_settings(self) -> None:
        self.settings.setValue("input", self.in_edit.text().strip())
        self.settings.setValue("output", self.out_edit.text().strip())
        self.settings.setValue("open", self.cb_open.isChecked())
        self.settings.setValue("log", self.cb_log.isChecked())
        self.settings.setValue("merge", self.cb_merge.isChecked())

    def _pick_input(self) -> None:
        d = QFileDialog.getExistingDirectory(self, "Chọn thư mục chứa Excel", self.in_edit.text() or str(Path.home()))
        if d:
            self.in_edit.setText(d)
            if not self.out_edit.text().strip():
                self.out_edit.setText(str(Path(d).parent / "dataCSV_UTF8"))

    def _pick_output(self) -> None:
        d = QFileDialog.getExistingDirectory(self, "Chọn thư mục ghi CSV", self.out_edit.text() or str(Path.home()))
        if d:
            self.out_edit.setText(d)

    def _open_output(self) -> None:
        p = self._last_output or Path(self.out_edit.text().strip())
        if p.is_dir():
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(p)))
        else:
            QMessageBox.information(self, "Thư mục", "Thư mục đầu ra chưa tồn tại.")

    def _clear_log(self) -> None:
        self.log_view.clear()
        self._log_lines.clear()
        self.table.setRowCount(0)

    def _append_log(self, level: str, message: str) -> None:
        colors = {"OK": "#3F7D28", "FAIL": "#C0392B", "WARN": "#B9770E", "INFO": "#1769AA"}
        color = colors.get(level, "#526476")
        stamp = datetime.now().strftime("%H:%M:%S")
        self._log_lines.append((stamp, level, message))
        safe = html.escape(message)
        self.log_view.append(
            f'<span style="color:#8593A1;">[{stamp}]</span> '
            f'<span style="color:{color}; font-weight:600;">{level:4}</span> '
            f'<span style="color:#0B2C4D;">{safe}</span>'
        )
        sb = self.log_view.verticalScrollBar()
        sb.setValue(sb.maximum())

    # -------------------------------------------------------------- actions
    def _start(self) -> None:
        if self.worker and self.worker.isRunning():
            return
        in_dir = Path(self.in_edit.text().strip())
        out_dir = Path(self.out_edit.text().strip())
        if not in_dir.is_dir():
            QMessageBox.warning(self, "Thiếu thư mục", f"Không tìm thấy thư mục đầu vào:\n{in_dir}")
            return
        if not str(out_dir).strip():
            QMessageBox.warning(self, "Thiếu thư mục", "Hãy chọn thư mục đầu ra.")
            return
        try:
            out_dir.mkdir(parents=True, exist_ok=True)
        except Exception as e:  # noqa: BLE001
            QMessageBox.critical(self, "Không tạo được thư mục", str(e))
            return
        existing = list(out_dir.glob("*.csv"))
        if existing:
            r = QMessageBox.question(
                self,
                "Thư mục đã có CSV",
                f"Thư mục đầu ra đang có {len(existing)} tệp CSV.\n"
                "Tiếp tục sẽ ghi đè các tệp trùng tên. Bạn có muốn tiếp tục?",
                QMessageBox.Yes | QMessageBox.No,
                QMessageBox.No,
            )
            if r != QMessageBox.Yes:
                return

        self._save_settings()
        self._clear_log()
        self._last_output = out_dir
        self.start_btn.setEnabled(False)
        self.stop_btn.setEnabled(True)
        self.progress.setRange(0, 100)
        self.progress.setValue(0)
        self.progress_lbl.setText("Đang chuẩn bị …")
        self.state_lbl.setText("Đang chạy")
        self.card_files.set_value("0 / 0")
        self.card_sheets.set_value("0")
        self.card_failed.set_value("0")
        self.card_time.set_value("0:00")

        self.worker = ConvertWorker(
            in_dir=in_dir,
            out_dir=out_dir,
            merge_alldata=self.cb_merge.isChecked(),
        )
        self.worker.log.connect(self._append_log)
        self.worker.progress.connect(self._on_progress)
        self.worker.file_done.connect(self._on_file_done)
        self.worker.done.connect(self._on_done)
        self.worker.start()

    def _stop(self) -> None:
        if self.worker and self.worker.isRunning():
            self.worker.request_stop()
            self.stop_btn.setEnabled(False)
            self.state_lbl.setText("Đang dừng …")
            self.progress_lbl.setText("Sẽ dừng sau tệp hiện tại …")

    def _on_progress(self, current: int, total: int, filename: str) -> None:
        self.progress.setRange(0, max(total, 1))
        self.progress.setValue(current)
        if filename:
            self.progress_lbl.setText(f"Đang xử lý {current + 1}/{total}: {filename}")
            self.card_files.set_value(f"{current} / {total}")

    def _on_file_done(self, filename: str, sheets: int, error: str) -> None:
        row = self.table.rowCount()
        self.table.insertRow(row)
        vals = [filename, str(sheets) if not error else "—", "Lỗi" if error else "OK", error or ""]
        for c, v in enumerate(vals):
            item = QTableWidgetItem(v)
            if c == 2:
                item.setForeground(QColor("#C0392B" if error else "#3F7D28"))
                f = item.font()
                f.setBold(True)
                item.setFont(f)
            self.table.setItem(row, c, item)
        self.table.scrollToBottom()

    def _on_done(self, stats: dict) -> None:
        self.start_btn.setEnabled(True)
        self.stop_btn.setEnabled(False)
        self.state_lbl.setText("Đã dừng" if stats.get("stopped") else "Hoàn tất")
        total = max(int(stats.get("files", 0)), 1)
        self.progress.setRange(0, total)
        self.progress.setValue(int(stats.get("files", 0)) if not stats.get("stopped") else self.progress.value())
        self.card_files.set_value(f"{stats.get('ok', 0)} / {stats.get('files', 0)}")
        self.card_sheets.set_value(f"{stats.get('sheets', 0):,}".replace(",", "."))
        self.card_failed.set_value(str(stats.get("failed", 0)))
        secs = int(round(stats.get("elapsed", 0)))
        self.card_time.set_value(f"{secs // 60}:{secs % 60:02d}")
        self.progress_lbl.setText(
            f"Xong: {stats.get('ok', 0)} tệp OK, {stats.get('failed', 0)} tệp lỗi, "
            f"{stats.get('sheets', 0)} sheet trong {secs}s"
            + ("  ·  đã gộp AllData.csv" if stats.get("merged") else "")
        )

        if self.cb_log.isChecked() and self._log_lines:
            try:
                out = Path(stats.get("output", ""))
                out.mkdir(parents=True, exist_ok=True)
                log_path = out / f"_log_{datetime.now():%Y%m%d_%H%M%S}.txt"
                with log_path.open("w", encoding="utf-8") as f:
                    for stamp, level, msg in self._log_lines:
                        f.write(f"[{stamp}] {level:4} {msg}\n")
                self._append_log("INFO", f"Đã ghi nhật ký: {log_path.name}")
            except Exception as e:  # noqa: BLE001
                self._append_log("WARN", f"Không ghi được log: {e}")

        if self.cb_open.isChecked():
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(stats.get("output", ""))))

        if stats.get("stopped"):
            QMessageBox.information(
                self,
                "Đã dừng",
                f"Đã dừng theo yêu cầu.\n"
                f"Đã xử lý {stats.get('ok', 0)} tệp, {stats.get('sheets', 0)} sheet.",
            )
        else:
            icon = QMessageBox.Warning if stats.get("failed") else QMessageBox.Information
            QMessageBox(icon, "Hoàn tất", self.progress_lbl.text(), QMessageBox.Ok, self).exec()

    def closeEvent(self, event):  # noqa: N802
        if self.worker and self.worker.isRunning():
            r = QMessageBox.question(
                self,
                "Đang chuyển đổi",
                "Tiến trình đang chạy. Dừng và thoát?",
                QMessageBox.Yes | QMessageBox.No,
                QMessageBox.No,
            )
            if r != QMessageBox.Yes:
                event.ignore()
                return
            self.worker.request_stop()
            self.worker.wait(4000)
        self._save_settings()
        event.accept()


def apply_style(app: QApplication) -> None:
    app.setStyle("Fusion")
    app.setStyleSheet(
        """
        * { font-family: "Segoe UI", Arial, sans-serif; font-size: 13px; color: #0B2C4D; }
        QMainWindow, QWidget#Central { background: #F3F7FB; }
        QFrame[card="true"] { background: #FFFFFF; border: 1px solid #D7E3EF; border-radius: 12px; }
        QLineEdit {
            background: #FBFDFF; border: 1px solid #C9D8E6; border-radius: 7px;
            padding: 8px 10px; selection-background-color: #1769AA;
        }
        QLineEdit:focus { border: 1px solid #1769AA; }
        QPushButton {
            background: #FFFFFF; border: 1px solid #C9D8E6; border-radius: 7px;
            padding: 8px 14px; color: #0B2C4D; font-weight: 600;
        }
        QPushButton:hover { background: #EAF2FA; }
        QPushButton:disabled { color: #A7B4C0; background: #F4F7FA; }
        QPushButton#Primary { background: #1769AA; border: 1px solid #1769AA; color: white; }
        QPushButton#Primary:hover { background: #10507F; border-color: #10507F; }
        QPushButton#Primary:disabled { background: #9DB8CF; border-color: #9DB8CF; }
        QPushButton#Danger { background: #FFFFFF; border: 1px solid #E8891C; color: #B9550B; }
        QPushButton#Danger:hover { background: #FDF2E4; }
        QPushButton#Danger:disabled { color: #C9B9A6; border-color: #E4D4C2; }
        QPushButton#Secondary { background: #F5F9FD; }
        QPushButton#Ghost { background: transparent; border: 1px solid transparent; color: #526476; font-weight: 500; }
        QPushButton#Ghost:hover { background: #EAF2FA; }
        QCheckBox { spacing: 7px; color: #526476; font-size: 12px; }
        QCheckBox::indicator { width: 16px; height: 16px; border: 1px solid #B9CBDD; border-radius: 4px; background: white; }
        QCheckBox::indicator:checked { background: #1769AA; border-color: #1769AA; }
        QProgressBar { border: none; background: #E6EDF4; border-radius: 5px; }
        QProgressBar::chunk { background: #5FA83C; border-radius: 5px; }
        QTabWidget::pane { border: 1px solid #D7E3EF; border-radius: 10px; background: #FFFFFF; top: -1px; }
        QTabBar::tab { background: transparent; padding: 8px 16px; color: #526476; border: none; }
        QTabBar::tab:selected { color: #1769AA; font-weight: 700; border-bottom: 2px solid #1769AA; }
        QTabBar::tab:hover { color: #1769AA; }
        QTableWidget { border: none; gridline-color: #E6EDF4; alternate-background-color: #F7FAFD; }
        QHeaderView::section {
            background: #EAF2FA; border: none; padding: 7px; font-weight: 700;
            color: #0B2C4D; font-size: 11px;
        }
        QScrollBar:vertical { background: transparent; width: 10px; margin: 2px; }
        QScrollBar::handle:vertical { background: #C9D8E6; border-radius: 5px; min-height: 24px; }
        QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical { height: 0; }
        QToolTip { background: #0B2C4D; color: white; border: none; padding: 5px; }
        """
    )


def run_selftest() -> int:
    """Kiểm tra lõi mà không mở GUI — dùng cho kiểm thử bản EXE."""
    import openpyxl

    with tempfile.TemporaryDirectory(prefix="x2c_gui_selftest_") as td:
        root = Path(td)
        inp, out = root / "in", root / "out"
        inp.mkdir()
        out.mkdir()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        ws["A1"] = "b¸o c¸o kÕt qu¶"
        ws["B1"] = "STT"
        ws.append(["", 1])
        wb.save(inp / "sample.xlsx")
        wb.close()
        n = core.convert_xlsx(inp / "sample.xlsx", out, set())
        files = list(out.glob("*.csv"))
        assert n == 1 and len(files) == 1, (n, files)
        text = files[0].read_text(encoding="utf-8-sig")
        assert "báo cáo kết quả" in text, text[:100]
    print("SELFTEST OK")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=f"{APP_NAME} {APP_VERSION}")
    ap.add_argument("--selftest", action="store_true", help="Kiểm tra lõi rồi thoát")
    ap.add_argument("--smoke-gui", action="store_true", help="Mở GUI nhanh rồi thoát")
    ap.add_argument("--cli", action="store_true", help="Chạy batch không mở GUI")
    ap.add_argument("--input", default=None)
    ap.add_argument("--output", default=None)
    ap.add_argument("--merge", action="store_true")
    args = ap.parse_args()

    if args.selftest:
        return run_selftest()

    if args.cli:
        argv = []
        if args.input:
            argv += ["--input", args.input]
        if args.output:
            argv += ["--output", args.output]
        rc = core.main(argv)
        if rc == 0 and args.merge:
            out_dir = (
                Path(args.output).resolve()
                if args.output
                else (Path(__file__).resolve().parent.parent / "dataCSV").resolve()
            )
            out_dir.mkdir(parents=True, exist_ok=True)
            try:
                import normalize

                rc = normalize.main(
                    [
                        "--input",
                        str(out_dir),
                        "--max-files",
                        "0",
                        "--output",
                        str(out_dir / "AllData.csv"),
                    ]
                )
            except Exception as e:  # noqa: BLE001
                print(f"Lỗi gộp AllData.csv: {e}", file=sys.stderr)
                rc = 1
        return rc

    app = QApplication(sys.argv)
    app.setApplicationName(APP_NAME)
    app.setOrganizationName(APP_ORG)
    icon_path = resource_path("assets/app.ico")
    if icon_path.exists():
        app.setWindowIcon(QIcon(str(icon_path)))
    apply_style(app)

    win = MainWindow()
    win.show()
    if args.smoke_gui:
        def _smoke_quit() -> None:
            win.close()
            app.quit()

        QTimer.singleShot(500, _smoke_quit)
        app.exec()
        # Thoát cứng trong chế độ kiểm thử để không để lại process onefile.
        os._exit(0)
    return app.exec()


if __name__ == "__main__":
    raise SystemExit(main())



