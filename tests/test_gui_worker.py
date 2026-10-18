# -*- coding: utf-8 -*-
"""Test luồng QThread của GUI (headless/offscreen)."""
from __future__ import annotations

import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))


class TestGuiWorker(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        try:
            from PySide6.QtWidgets import QApplication
        except Exception as e:  # noqa: BLE001
            raise unittest.SkipTest(f"Không có PySide6: {e}")
        cls.app = QApplication.instance() or QApplication([])

    def test_worker_emits_stats(self) -> None:
        import openpyxl
        from PySide6.QtCore import QEventLoop, QTimer

        import gui

        td = Path(tempfile.mkdtemp(prefix="x2c_gui_"))
        try:
            inp, out = td / "in", td / "out"
            inp.mkdir()
            out.mkdir()
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "S1"
            ws["A1"] = "b¸o c¸o"
            ws["B2"] = 1
            wb.save(inp / "f1.xlsx")
            wb.close()

            worker = gui.ConvertWorker(inp, out, merge_alldata=False)
            done: dict = {}
            logs: list[tuple[str, str]] = []
            worker.log.connect(lambda lvl, msg: logs.append((lvl, msg)))
            worker.done.connect(lambda d: done.update(d))
            loop = QEventLoop()
            worker.done.connect(lambda _d: loop.quit())
            worker.start()
            QTimer.singleShot(30000, loop.quit)
            loop.exec()
            worker.wait(3000)

            self.assertEqual(done.get("ok"), 1)
            self.assertEqual(done.get("sheets"), 1)
            self.assertEqual(done.get("failed"), 0)
            self.assertEqual(len(list(out.glob("*.csv"))), 1)
            self.assertTrue(any(lvl == "OK" for lvl, _ in logs))
        finally:
            shutil.rmtree(td, ignore_errors=True)


if __name__ == "__main__":
    unittest.main(verbosity=2)

