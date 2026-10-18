# -*- coding: utf-8 -*-
"""Test nhanh lõi Xlsx2CSV / Normalize (chạy: python -m unittest discover -s tests)."""
from __future__ import annotations

import csv
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

import xlsx2csv as core  # noqa: E402


class TestTcvn3(unittest.TestCase):
    def test_convert_common_sentence(self) -> None:
        src = "b¸o c¸o kÕt qu¶ thÝ nghiÖm ®Êt"
        self.assertTrue(core.looks_like_tcvn3(src))
        self.assertEqual(core.tcvn3_to_unicode(src), "báo cáo kết quả thí nghiệm đất")

    def test_unicode_text_is_not_mistaken(self) -> None:
        # "núi" có ú là chữ Unicode hợp lệ, không được coi là TCVN3.
        self.assertFalse(core.looks_like_tcvn3("núi đất"))
        self.assertFalse(core.looks_like_tcvn3("báo cáo kết quả"))
        # micro-mét không phải TCVN3
        self.assertFalse(core.looks_like_tcvn3("hạt 10 µm"))

    def test_font_detection(self) -> None:
        self.assertTrue(core.is_tcvn3_font(".VnTime"))
        self.assertTrue(core.is_tcvn3_font("TCVN3"))
        self.assertFalse(core.is_tcvn3_font("Arial"))

    def test_safe_filename(self) -> None:
        self.assertEqual(core.safe_filename_piece("  a/b:c*?  "), "a_b_c_")
        self.assertEqual(core.safe_filename_piece("   "), "sheet")


class TestAngles(unittest.TestCase):
    def test_ddmm_to_decimal(self) -> None:
        self.assertAlmostEqual(core._convert_angle_numeric(708), 7 + 8 / 60, places=8)
        self.assertAlmostEqual(core._convert_angle_numeric(1642), 16 + 42 / 60, places=8)

    def test_dms_text(self) -> None:
        self.assertAlmostEqual(
            core._convert_angle_cell_for_csv("19°17′"), 19 + 17 / 60, places=8
        )


class TestConvertXlsx(unittest.TestCase):
    def setUp(self) -> None:
        import openpyxl

        self.tmp = Path(tempfile.mkdtemp(prefix="x2c_test_"))
        self.in_dir = self.tmp / "in"
        self.out_dir = self.tmp / "out"
        self.in_dir.mkdir()
        self.out_dir.mkdir()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        a1 = ws["A1"]
        a1.value = "b¸o c¸o kÕt qu¶"
        a1.font = openpyxl.styles.Font(name=".VnTime")
        ws["B1"] = "STT"
        ws["C1"] = "Mẫu"
        ws["D1"] = "Hố"
        ws["E1"] = "Từ"
        ws["F1"] = "Đến"
        ws["G1"] = "Góc ma sát trong"
        ws.append(["", 1, "M1", "H1", 1, 2, 708])
        ws2 = wb.create_sheet("Empty")
        ws2["A1"] = "   "
        self.xlsx = self.in_dir / "sample.xlsx"
        wb.save(self.xlsx)
        wb.close()

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_convert_and_angle(self) -> None:
        used: set[str] = set()
        n = core.convert_xlsx(self.xlsx, self.out_dir, used)
        self.assertEqual(n, 1)
        outs = list(self.out_dir.glob("*.csv"))
        self.assertEqual(len(outs), 1)
        self.assertEqual(outs[0].name, "sample_Data.csv")
        raw = outs[0].read_bytes()
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))
        with outs[0].open("r", encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))
        self.assertEqual(rows[0][0], "báo cáo kết quả")
        self.assertAlmostEqual(float(rows[1][6]), 7 + 8 / 60, places=6)

    def test_same_stem_does_not_overwrite(self) -> None:
        used: set[str] = set()
        core.convert_xlsx(self.xlsx, self.out_dir, used)
        core.convert_xlsx(self.xlsx, self.out_dir, used)
        names = sorted(p.name for p in self.out_dir.glob("*.csv"))
        self.assertEqual(names, ["sample_Data.csv", "sample_Data_2.csv"])

    def test_discover(self) -> None:
        found = core.discover_workbooks(self.in_dir)
        self.assertEqual(found, [self.xlsx.resolve()])


class TestConvertXls(unittest.TestCase):
    def test_real_xls(self) -> None:
        src = ROOT / "examples" / "sample.xls"
        if not src.exists():
            self.skipTest("Không có examples/sample.xls")
        with tempfile.TemporaryDirectory(prefix="x2c_xls_") as td:
            out = Path(td)
            n = core.convert_xls(src, out, set())
            self.assertGreaterEqual(n, 1)
            self.assertGreaterEqual(len(list(out.glob("*.csv"))), 1)


class TestNormalize(unittest.TestCase):
    def test_normalize_angle_and_row(self) -> None:
        import normalize as norm

        self.assertEqual(norm.normalize_angle_cell_value("708"), "7.13333333")
        self.assertEqual(norm.normalize_angle_cell_value("19°17′"), "19.28333333")
        self.assertTrue(norm.looks_like_data_row(["", "1", "M1", "H1", "1", "2", "x"]))
        self.assertFalse(norm.looks_like_data_row(["", "", "M1", "H1", "1", "2", "x"]))


if __name__ == "__main__":
    unittest.main(verbosity=2)


