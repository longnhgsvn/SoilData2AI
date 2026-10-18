# -*- coding: utf-8 -*-
"""Tạo bộ dữ liệu mẫu nhỏ (không chứa dữ liệu thật) cho ví dụ và test."""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font

HERE = Path(__file__).resolve().parent


def main() -> int:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"

    a1 = ws["A1"]
    a1.value = "b¸o c¸o kÕt qu¶ thÝ nghiÖm ®Êt"
    a1.font = Font(name=".VnTime", size=12)

    for cell, value in {
        "B1": "STT",
        "C1": "Số hiệu mẫu",
        "D1": "Hố khoan",
        "E1": "Từ (m)",
        "F1": "Đến (m)",
        "G1": "Góc ma sát trong",
        "H1": "Độ ẩm (%)",
    }.items():
        ws[cell] = value

    ws.append(["", 1, "M1", "HK1", 1.0, 2.0, 708, 22.5])
    ws.append(["", 2, "M2", "HK1", 2.0, 3.0, 1642, 24.1])

    empty = wb.create_sheet("Empty")
    empty["A1"] = "   "
    wb.save(HERE / "sample.xlsx")
    wb.close()
    print(f"OK: {HERE / 'sample.xlsx'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
